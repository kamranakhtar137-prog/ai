# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ---------- Styles ----------
HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
SUB_FONT = Font(italic=True, size=10, color="666666")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_COLORS = {
    "Done": "C6EFCE",
    "Partial": "FFEB9C",
    "Broken": "FFC7CE",
    "Not Started": "FFC7CE",
    "Not Yet Audited": "D9D9D9",
}
STATUS_FONT_COLORS = {
    "Done": "006100",
    "Partial": "9C6500",
    "Broken": "9C0006",
    "Not Started": "9C0006",
    "Not Yet Audited": "595959",
}

def status_fill(status):
    return PatternFill(start_color=STATUS_COLORS.get(status, "FFFFFF"),
                        end_color=STATUS_COLORS.get(status, "FFFFFF"), fill_type="solid")

def status_font(status):
    return Font(bold=True, color=STATUS_FONT_COLORS.get(status, "000000"), size=10)

# ======================================================================
# SHEET 1: Executive Summary
# ======================================================================
ws = wb.active
ws.title = "Executive Summary"
ws.sheet_view.showGridLines = False

ws["B2"] = "Main Street Collectibles — On-Page SEO Audit & Tracker"
ws["B2"].font = TITLE_FONT
ws["B3"] = "Live-site verification as of July 27, 2026 (mainstreetnewcanaan.com)"
ws["B3"].font = SUB_FONT

summary_rows = [
    ("Page", "Status", "Key Issue"),
    ("Home ( / )", "Partial", "Title fixed. Meta description is bloated (226 chars, has leftover text). H1 is a hidden logo tag only, no visible keyword H1. Planned H2 sections (Shop by Category, Why Shop With Us, Visit Our Store) not built."),
    ("Shop/Catalog ( /collections/trading-card )", "Partial", "Title + H1 fixed. Meta description has a text-duplication bug (\"baBrowse oseball\"). Planned H2 category sections not built."),
    ("Baseball ( /collections/baseball )", "Broken", "URL returns 404 (collection unpublished, no products). None of the planned title/description/H1/H2 are live."),
    ("Contact ( /pages/contact )", "Partial", "Title + description fixed and correct. No H1 tag exists on the page at all. Planned H2 sections not built."),
    ("Hours & Info ( /pages/hours-and-info )", "Partial", "Title + description fixed. H1 live as \"Hours and Info\" instead of planned \"Store Hours & Location\". Planned H2 sections not built."),
    ("TCG / Pokémon", "Not Yet Audited", "Not checked in this pass."),
    ("Basketball / WNBA", "Not Yet Audited", "Not checked in this pass."),
    ("Football / NFL", "Not Yet Audited", "No dedicated page planned yet — pending client scope decision."),
    ("Hockey / NHL", "Not Yet Audited", "No dedicated page planned yet."),
    ("Merchandise (product template)", "Not Yet Audited", "Not checked in this pass."),
    ("About", "Not Yet Audited", "URL decision (/pages/info vs /pages/about) still needs confirmation."),
    ("Gift Cards", "Not Yet Audited", "Not checked in this pass."),
    ("Product page template", "Not Yet Audited", "Not checked in this pass."),
]

start_row = 5
for i, row in enumerate(summary_rows):
    r = start_row + i
    for j, val in enumerate(row):
        c = ws.cell(row=r, column=2 + j, value=val)
        c.border = BORDER
        c.alignment = WRAP
        if i == 0:
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
        elif j == 1:
            c.fill = status_fill(val)
            c.font = status_font(val)

ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 80
for i in range(len(summary_rows)):
    ws.row_dimensions[start_row + i].height = 34
ws.row_dimensions[start_row].height = 18

# Priority punch list
punch_row = start_row + len(summary_rows) + 2
ws.cell(row=punch_row, column=2, value="Priority Fix List (in order)").font = Font(bold=True, size=12, color="1F3864")
punch_items = [
    "1. Baseball collection is 404ing publicly — either add product(s) so the collection can be published, or intentionally unpublish/redirect it so it doesn't sit as a dead link in the sitemap.",
    "2. Fix the corrupted Shop page meta description (\"baBrowse oseball\" text-duplication bug) — likely a copy/paste error in the Shopify SEO field.",
    "3. Trim the Home page meta description back down to ~155-160 chars — it currently has extra trailing text pushing it to 226 chars, which Google will cut off mid-sentence.",
    "4. Add a real visible H1 to the Contact page — it currently has none.",
    "5. Align the Hours & Info H1 to \"Store Hours & Location\" (currently \"Hours and Info\").",
    "6. Add a real, visible, keyword-rich H1 to Home (currently only a visually-hidden logo H1) and build out the three planned H2 content sections.",
    "7. Build the four planned H2 category sections on the Shop page (Pokémon Cards / Sports Cards / Sealed Product / Graded & Vintage Singles).",
    "8. Build the three planned H2 sections on Contact and on Hours & Info.",
    "9. Audit and fix the remaining not-yet-checked pages (TCG/Pokémon, Basketball/WNBA, Merchandise, About, Gift Cards, product template) using the same method.",
]
for i, item in enumerate(punch_items):
    ws.cell(row=punch_row + 1 + i, column=2, value=item).alignment = WRAP
    ws.merge_cells(start_row=punch_row + 1 + i, start_column=2, end_row=punch_row + 1 + i, end_column=4)
    ws.row_dimensions[punch_row + 1 + i].height = 28

# ======================================================================
# SHEET 2: Full Page Tracker (planned vs live)
# ======================================================================
ws2 = wb.create_sheet("Page Tracker")
ws2.sheet_view.showGridLines = False

headers = [
    "Page", "URL", "Meta Title (Planned)", "Meta Title (Live)", "Title Status",
    "Meta Description (Planned)", "Meta Description (Live)", "Description Status",
    "H1 (Planned)", "H1 (Live)", "H1 Status",
    "H2/H3 (Planned)", "H2 (Live)", "H2 Status",
    "Overall Status", "Notes / Next Action",
]
for j, h in enumerate(headers, start=1):
    c = ws2.cell(row=1, column=j, value=h)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = CENTER
    c.border = BORDER
ws2.freeze_panes = "A2"

data = [
    dict(
        page="Home", url="/",
        title_planned="Hobby Shop in New Canaan, CT | Main Street Collectibles",
        title_live="Hobby Shop in New Canaan, CT | Main Street Collectibles",
        title_status="Done",
        desc_planned="Your one-stop shop to buy and sell trading cards, sports cards & memorabilia. Visit us in-person or shop online: Pokémon, baseball, basketball cards + more.",
        desc_live="Buy and sell trading cards, sports cards & memorabilia at our New Canaan, CT hobby shop. Shop online or in-person: Pokémon, baseball, basketball cards + more. New Canaan's premier destination for trading cards and collectibles (226 chars — too long, trim it)",
        desc_status="Partial",
        h1_planned="Trading Cards & Sports Collectibles in New Canaan, CT",
        h1_live="\"Main Street Collectibles\" (visually-hidden logo H1 only — no visible keyword H1 exists)",
        h1_status="Not Started",
        h2_planned="Shop by Category (Pokemon/Baseball/Basketball/TCG); Why Shop With Us; Visit Our New Canaan Store",
        h2_live="None of the planned sections exist. Only theme chrome (cart drawer, newsletter signup, search)",
        h2_status="Not Started",
        overall="Partial",
        notes="Title is correct. Fix bloated meta description. Add a real content section with a visible H1, then add the 3 planned H2 sections.",
    ),
    dict(
        page="Shop / Catalog", url="/collections/trading-card (redirected from /collections/online-store — note: singular \"trading-card\", plan said plural \"trading-cards\")",
        title_planned="Trading Cards for Sale | Main Street Collectibles New Canaan",
        title_live="Trading Cards for Sale | Main Street Collectibles New Canaan",
        title_status="Done",
        desc_planned="Browse our trading card collection online or visit our hobby shop in New Canaan, CT. Pokémon, baseball & basketball cards, sealed boxes, graded singles!",
        desc_live="BUG: \"...Pokémon, baBrowse oseball & basketball cards...\" — a text-duplication/copy-paste error has corrupted the word \"baseball\" in the live field.",
        desc_status="Broken",
        h1_planned="Shop Trading & Sports Cards",
        h1_live="Shop Trading & Sports Cards",
        h1_status="Done",
        h2_planned="Pokemon Cards; Sports Cards (Baseball & Basketball); Sealed Product & Booster Boxes; Graded & Vintage Singles",
        h2_live="None of the planned sections exist. Only theme chrome (cart drawer, filter drawer, newsletter, search)",
        h2_status="Not Started",
        overall="Partial",
        notes="301 redirect from old URL confirmed working. Fix the corrupted meta description immediately. Build the 4 category H2 sections. Confirm whether URL should be plural per original plan.",
    ),
    dict(
        page="Baseball", url="/collections/baseball",
        title_planned="Baseball Cards for Sale | Main Street Collectibles, CT",
        title_live="404 Not Found – Main Street Collectibles",
        title_status="Broken",
        desc_planned="Browse incredible baseball card collection online or visit our hobby shop in New Canaan, CT. Vintage packs, rookie cards, graded singles + more!",
        desc_live="(none — page 404s)",
        desc_status="Broken",
        h1_planned="Baseball Cards",
        h1_live="Page not found",
        h1_status="Broken",
        h2_planned="Vintage & Rookie Baseball Cards; Baseball Packs & Sealed Product; Sell Your Baseball Cards",
        h2_live="(none — page 404s)",
        h2_status="Broken",
        overall="Broken",
        notes="Collection is unpublished/empty (no products), so it 404s publicly. Either add product(s) and publish, or leave unpublished until inventory exists — but do not leave it linked/indexed as a live 404.",
    ),
    dict(
        page="Contact", url="/pages/contact",
        title_planned="Contact Main Street Collectibles | New Canaan, CT",
        title_live="Contact Main Street Collectibles | New Canaan, CT",
        title_status="Done",
        desc_planned="Come visit our hobby shop in New Canaan, CT! Main Street Collectibles address, hours, directions, and info on how to buy or sell cards.",
        desc_live="Come visit our hobby shop in New Canaan, CT! Main Street Collectibles address, hours, directions, and info on how to buy or sell cards.",
        desc_status="Done",
        h1_planned="Contact Main Street Collectibles",
        h1_live="(none — no H1 tag exists anywhere on the page)",
        h1_status="Not Started",
        h2_planned="Visit Our Store; Sell Your Cards; Store Hours & Directions",
        h2_live="None of the planned sections exist. Only theme chrome (cart drawer, newsletter, search)",
        h2_status="Not Started",
        overall="Partial",
        notes="Title and description are correct and live. Add the missing H1 and build the 3 planned H2 sections.",
    ),
    dict(
        page="Hours & Info", url="/pages/hours-and-info",
        title_planned="Store Hours & Location | Main Street Collectibles New Canaan",
        title_live="Store Hours & Location | Main Street Collectibles New Canaan",
        title_status="Done",
        desc_planned="Come visit our hobby shop in New Canaan, CT! Main Street Collectibles hours, location, and directions.",
        desc_live="Come visit our hobby shop in New Canaan, CT! Main Street Collectibles hours, location, and directions.",
        desc_status="Done",
        h1_planned="Store Hours & Location",
        h1_live="Hours and Info",
        h1_status="Partial",
        h2_planned="Today's Hours; Directions & Parking; Trade Nights / Events (if applicable)",
        h2_live="None of the planned sections exist. Only theme chrome (cart drawer, newsletter, search)",
        h2_status="Not Started",
        overall="Partial",
        notes="Title and description are correct and live. Update H1 text to exactly match the plan, then build the 3 planned H2 sections.",
    ),
    dict(
        page="TCG / Pokémon", url="/collections/tcg-card",
        title_planned="Pokémon Cards for Sale | Main Street Collectibles New Canaan",
        title_live="Not yet audited", title_status="Not Yet Audited",
        desc_planned="Browse our incredible collection online or visit our hobby shop in New Canaan, CT and consult with our expert team. Rare singles, packs, sealed product + more!",
        desc_live="Not yet audited", desc_status="Not Yet Audited",
        h1_planned="Pokémon TCG Cards", h1_live="Not yet audited", h1_status="Not Yet Audited",
        h2_planned="Pokemon Singles & Sealed Product; Vintage & 1st Edition Pokemon; Other TCG (MTG, Yu-Gi-Oh & More)",
        h2_live="Not yet audited", h2_status="Not Yet Audited",
        overall="Not Yet Audited", notes="Check this page live using the same method as the 5 pages above.",
    ),
    dict(
        page="Basketball / WNBA", url="/collections/basketball, /collections/wnba",
        title_planned="No dedicated title — folded into Shop page (needs decision)",
        title_live="Not yet audited", title_status="Not Yet Audited",
        desc_planned="No dedicated description — covered by Shop page for now (needs decision)",
        desc_live="Not yet audited", desc_status="Not Yet Audited",
        h1_planned="N/A — not a standalone page at this time", h1_live="Not yet audited", h1_status="Not Yet Audited",
        h2_planned="N/A", h2_live="Not yet audited", h2_status="Not Yet Audited",
        overall="Not Yet Audited",
        notes="These URLs already exist live — decide whether to give them unique metadata or canonicalize/noindex them to the Shop page so they don't sit as duplicate/thin content.",
    ),
    dict(
        page="Football / NFL", url="No dedicated collection — surfaces via Shop only",
        title_planned="No dedicated title — pending client scope decision",
        title_live="N/A", title_status="Not Yet Audited",
        desc_planned="Covered by Shop page for now",
        desc_live="N/A", desc_status="Not Yet Audited",
        h1_planned="N/A — not a standalone page at this time", h1_live="N/A", h1_status="Not Yet Audited",
        h2_planned="N/A", h2_live="N/A", h2_status="Not Yet Audited",
        overall="Not Yet Audited",
        notes="New finding — needs client decision on whether this becomes a standalone collection page.",
    ),
    dict(
        page="Hockey / NHL", url="No dedicated collection — surfaces via Shop only",
        title_planned="No dedicated title — folded into Shop page",
        title_live="N/A", title_status="Not Yet Audited",
        desc_planned="Covered by Shop page for now",
        desc_live="N/A", desc_status="Not Yet Audited",
        h1_planned="N/A — not a standalone page at this time", h1_live="N/A", h1_status="Not Yet Audited",
        h2_planned="N/A", h2_live="N/A", h2_status="Not Yet Audited",
        overall="Not Yet Audited",
        notes="New finding — needs client decision on whether this becomes a standalone collection page.",
    ),
    dict(
        page="Merchandise (template)", url="/products/msc-trucker-hat (example)",
        title_planned="[Item Name] | Main Street Collectibles, New Canaan",
        title_live="Not yet audited", title_status="Not Yet Audited",
        desc_planned="Shop Main Street Collectibles branded merchandise, including [item name] and more! Shop online or visit our hobby shop in New Canaan, CT.",
        desc_live="Not yet audited", desc_status="Not Yet Audited",
        h1_planned="[Item Name]", h1_live="Not yet audited", h1_status="Not Yet Audited",
        h2_planned="N/A — standard product page, no card-specific sections needed",
        h2_live="Not yet audited", h2_status="Not Yet Audited",
        overall="Not Yet Audited",
        notes="Note: this exact item (MSC Trucker Hat) was found mixed into the Shop/Catalog collection grid during this audit — confirm it's tagged/categorized correctly and not duplicated across collections.",
    ),
    dict(
        page="About", url="/pages/info (CONFIRM: proposed /pages/about)",
        title_planned="About Main Street Collectibles | New Canaan CT Card Shop",
        title_live="Not yet audited", title_status="Not Yet Audited",
        desc_planned="Main Street Collectibles is a hobby shop in New Canaan, CT owned by passionate experts in trading cards and sports memorabilia. Read our story.",
        desc_live="Not yet audited", desc_status="Not Yet Audited",
        h1_planned="About Main Street Collectibles", h1_live="Not yet audited", h1_status="Not Yet Audited",
        h2_planned="Our Story; What We Buy & Sell; Why Collectors Shop With Us",
        h2_live="Not yet audited", h2_status="Not Yet Audited",
        overall="Not Yet Audited",
        notes="URL decision still open. If moving to /pages/about, set up a 301 redirect from /pages/info first, same as was done for the Shop page.",
    ),
    dict(
        page="Gift Cards", url="/products/msc-gift-card",
        title_planned="Gift Cards | Main Street Collectibles New Canaan",
        title_live="Not yet audited", title_status="Not Yet Audited",
        desc_planned="The perfect gift for anyone passionate about the hobby. We sell trading cards, sports cards & memorabilia. Shop online or visit our store in New Canaan, CT.",
        desc_live="Not yet audited", desc_status="Not Yet Audited",
        h1_planned="Main Street Collectibles Gift Cards", h1_live="Not yet audited", h1_status="Not Yet Audited",
        h2_planned="How Gift Cards Work; Redeeming In-Store or Online",
        h2_live="Not yet audited", h2_status="Not Yet Audited",
        overall="Not Yet Audited", notes="Check this page live using the same method as the 5 pages above.",
    ),
    dict(
        page="Product page (template)", url="/products/[handle]",
        title_planned="[Year] [Set] [Card Name] [Grade/Edition] | Main Street Collectibles",
        title_live="Not yet audited", title_status="Not Yet Audited",
        desc_planned="Buy the [Card Name], graded [condition/grade] with [key attribute]. In stock at Main Street Collectibles, New Canaan CT. [Price if available].",
        desc_live="Not yet audited", desc_status="Not Yet Audited",
        h1_planned="[Year] [Set] [Card Name] [Grade/Edition]", h1_live="Not yet audited", h1_status="Not Yet Audited",
        h2_planned="Card Details; Why Choose This Card; Authenticity & Condition; FAQ",
        h2_live="Not yet audited", h2_status="Not Yet Audited",
        overall="Not Yet Audited",
        notes="Needs a bulk-template approach (metafields/CSV import or SEO app) since this applies per-product, not a one-off manual edit.",
    ),
]

r = 2
for row in data:
    vals = [
        row["page"], row["url"], row["title_planned"], row["title_live"], row["title_status"],
        row["desc_planned"], row["desc_live"], row["desc_status"],
        row["h1_planned"], row["h1_live"], row["h1_status"],
        row["h2_planned"], row["h2_live"], row["h2_status"],
        row["overall"], row["notes"],
    ]
    for j, v in enumerate(vals, start=1):
        c = ws2.cell(row=r, column=j, value=v)
        c.border = BORDER
        c.alignment = WRAP
        header = headers[j - 1]
        if header.endswith("Status") or header == "Overall Status":
            c.fill = status_fill(v)
            c.font = status_font(v)
    ws2.row_dimensions[r].height = 60
    r += 1

widths = [20, 24, 26, 26, 12, 30, 30, 13, 22, 26, 12, 30, 26, 12, 14, 34]
for j, w in enumerate(widths, start=1):
    ws2.column_dimensions[get_column_letter(j)].width = w

ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{r-1}"

wb.save("/workspace/seo-audit/MSC_SEO_Audit_Tracker.xlsx")
print("Saved xlsx")
