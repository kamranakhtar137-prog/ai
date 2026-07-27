# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def shade_cell(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)

STATUS_COLORS = {
    "Done": "C6EFCE",
    "Partial": "FFEB9C",
    "Broken": "FFC7CE",
    "Not Started": "FFC7CE",
}

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)

title = doc.add_heading("Main Street Collectibles — On-Page SEO Fix Report", level=0)
sub = doc.add_paragraph()
sub_run = sub.add_run("Live-site verification of EOD status update — mainstreetnewcanaan.com — July 27, 2026")
sub_run.italic = True
sub_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_heading("Executive Summary", level=1)
doc.add_paragraph(
    "Five pages were reported as complete in the EOD update: Home, Shop/Catalog, Baseball, "
    "Contact, and Hours & Info. Each was checked directly against the live site's HTML "
    "(title tag, meta description, H1, and H2 headings). Findings:"
)

bullets = [
    "2 of 5 pages (Contact, Hours & Info) have correct title tags and meta descriptions live.",
    "All 5 pages are still missing the planned H2 content sections — the theme currently renders "
    "only default chrome (cart drawer, filter drawer, newsletter signup, search) with no custom "
    "on-page sections built yet.",
    "3 real bugs need fixing immediately: a corrupted/duplicated meta description on the Shop page, "
    "an oversized meta description on Home, and a missing H1 on Contact.",
    "The Baseball collection is currently returning a 404 to the public because it has no products "
    "and is unpublished — none of its planned SEO content is live.",
    "8 of the 13 planned pages/templates (TCG/Pokémon, Basketball/WNBA, Football, Hockey, "
    "Merchandise, About, Gift Cards, product template) have not yet been checked in this pass.",
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')

doc.add_heading("Status by Page", level=1)

pages = [
    dict(
        name="Home ( / )",
        overall="Partial",
        findings=[
            ("Meta Title", "Done", "Live and correct: \u201cHobby Shop in New Canaan, CT | Main Street Collectibles\u201d (55 chars)."),
            ("Meta Description", "Partial", "Live text is 226 characters (should be ~155\u2013160) with an extra trailing sentence appended after the intended copy: \u201c...basketball cards + more. New Canaan's premier destination for trading cards and collectibles\u201d. Google will truncate this mid-sentence. Needs to be trimmed back to the 157-char version."),
            ("H1", "Not Started", "Only H1 present is a visually-hidden logo tag (\u201cMain Street Collectibles\u201d) \u2014 not visible to users and not the planned keyword-rich heading. No real H1 content exists on the page."),
            ("H2 Sections", "Not Started", "None of the three planned sections (Shop by Category, Why Shop With Us, Visit Our New Canaan Store) have been built. The homepage template currently has no custom content sections at all."),
        ],
    ),
    dict(
        name="Shop / Catalog ( /collections/trading-card )",
        overall="Partial",
        findings=[
            ("URL", "Done", "301 redirect confirmed working from the old URL (/collections/online-store) to the new one. Note: live handle is singular \u201ctrading-card\u201d; original plan specified plural \u201ctrading-cards\u201d \u2014 confirm this is intentional."),
            ("Meta Title", "Done", "Live and correct: \u201cTrading Cards for Sale | Main Street Collectibles New Canaan\u201d (60 chars)."),
            ("Meta Description", "Broken", "Contains a text-duplication bug: the live text reads \u201c...Pok\u00e9mon, baBrowse oseball & basketball cards...\u201d \u2014 the word \u201cbaseball\u201d has been corrupted by an inserted duplicate fragment (\u201cBrowse o\u201d), almost certainly a copy/paste error in the Shopify SEO field. Needs to be replaced with the clean version."),
            ("H1", "Done", "Live and correct: \u201cShop Trading & Sports Cards\u201d."),
            ("H2 Sections", "Not Started", "None of the four planned category sections (Pokemon Cards, Sports Cards, Sealed Product & Booster Boxes, Graded & Vintage Singles) exist yet \u2014 page is a single flat product grid."),
        ],
    ),
    dict(
        name="Baseball ( /collections/baseball )",
        overall="Broken",
        findings=[
            ("Page Status", "Broken", "Returns HTTP 404 publicly. Confirmed cause: the collection has no products and \u201cShow in online store\u201d is disabled so it doesn't render an empty page \u2014 but this means the URL is currently dead to visitors and search engines."),
            ("Meta Title / Description / H1 / H2", "Broken", "None of the planned content is live since the page doesn't render at all (title tag shows Shopify's generic \u201c404 Not Found\u201d)."),
        ],
    ),
    dict(
        name="Contact ( /pages/contact )",
        overall="Partial",
        findings=[
            ("Meta Title", "Done", "Live and correct: \u201cContact Main Street Collectibles | New Canaan, CT\u201d (49 chars)."),
            ("Meta Description", "Done", "Live and correct, exact match to plan (135 chars)."),
            ("H1", "Not Started", "No H1 tag exists anywhere on this page \u2014 not even a hidden fallback like Home has. This is a real gap that should be fixed."),
            ("H2 Sections", "Not Started", "None of the three planned sections (Visit Our Store, Sell Your Cards, Store Hours & Directions) exist yet."),
        ],
    ),
    dict(
        name="Hours & Info ( /pages/hours-and-info )",
        overall="Partial",
        findings=[
            ("Meta Title", "Done", "Live and correct: \u201cStore Hours & Location | Main Street Collectibles New Canaan\u201d (60 chars)."),
            ("Meta Description", "Done", "Live and correct, exact match to plan (103 chars)."),
            ("H1", "Partial", "Live H1 reads \u201cHours and Info\u201d instead of the planned \u201cStore Hours & Location.\u201d Close, but not an exact match \u2014 should be updated for keyword consistency with the title tag."),
            ("H2 Sections", "Not Started", "None of the three planned sections (Today's Hours, Directions & Parking, Trade Nights / Events) exist yet."),
        ],
    ),
]

for pg in pages:
    doc.add_heading(pg["name"], level=2)
    p = doc.add_paragraph()
    r = p.add_run(f"Overall status: {pg['overall']}")
    r.bold = True

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    hdr[0].text = "Element"
    hdr[1].text = "Status"
    hdr[2].text = "Finding"
    for c in hdr:
        for para in c.paragraphs:
            for run in para.runs:
                run.bold = True

    for element, status, finding in pg["findings"]:
        row = table.add_row().cells
        row[0].text = element
        row[1].text = status
        row[2].text = finding
        shade_cell(row[1], STATUS_COLORS.get(status, "FFFFFF"))

    table.columns[0].width = Inches(1.3)
    table.columns[1].width = Inches(0.9)
    table.columns[2].width = Inches(4.3)
    doc.add_paragraph()

doc.add_heading("Priority Fix List", level=1)
fixes = [
    "Add at least one product to the Baseball collection (or re-enable \u201cShow in online store\u201d once inventory exists) so it stops returning 404, then apply the planned title/description/H1/H2.",
    "Replace the corrupted Shop page meta description \u2014 remove the duplicated \u201cBrowse o\u201d fragment from the middle of \u201cbaseball.\u201d",
    "Trim the Home page meta description back to ~155\u2013160 characters by removing the trailing sentence that was appended after the original approved copy.",
    "Add a visible H1 to the Contact page.",
    "Update the Hours & Info H1 from \u201cHours and Info\u201d to \u201cStore Hours & Location\u201d to match the plan.",
    "Add a real, visible, keyword-rich H1 to the Home page (currently only a hidden logo tag exists) and build its three planned H2 sections.",
    "Build the four planned H2 category sections on the Shop page.",
    "Build the three planned H2 sections each on Contact and on Hours & Info.",
    "Confirm the Shop URL should be singular (\u201c/collections/trading-card\u201d, as shipped) vs. the plural originally planned (\u201c/collections/trading-cards\u201d) \u2014 update the plan or the URL so they match.",
    "Once the above are done, run the same live-HTML audit on the remaining 8 pages/templates: TCG/Pok\u00e9mon, Basketball/WNBA, Football, Hockey, Merchandise, About, Gift Cards, and the product page template.",
]
for f in fixes:
    doc.add_paragraph(f, style='List Number')

doc.add_paragraph()
foot = doc.add_paragraph()
foot_run = foot.add_run("Companion file: MSC_SEO_Audit_Tracker.xlsx (full planned-vs-live tracker for all 13 pages/templates).")
foot_run.italic = True
foot_run.font.size = Pt(9)
foot_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.save("/workspace/seo-audit/MSC_SEO_Fix_Report.docx")
print("Saved docx")
